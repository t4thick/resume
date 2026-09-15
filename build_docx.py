from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import copy

doc = Document()

# ---------- Page margins ----------
for section in doc.sections:
    section.top_margin    = Inches(0.5)
    section.bottom_margin = Inches(0.5)
    section.left_margin   = Inches(0.75)
    section.right_margin  = Inches(0.75)

# ---------- Remove default paragraph spacing ----------
style = doc.styles['Normal']
style.font.name = 'Garamond'
style.font.size = Pt(10)
from docx.oxml.ns import qn
pPr = style.paragraph_format._element.get_or_add_pPr()
spacing = OxmlElement('w:spacing')
spacing.set(qn('w:before'), '0')
spacing.set(qn('w:after'),  '0')
pPr.append(spacing)

def set_spacing(para, before=0, after=0, line=None):
    pf = para.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after  = Pt(after)
    if line:
        from docx.shared import Pt as P
        pf.line_spacing = P(line)

def add_run(para, text, bold=False, italic=False, size=10, color=None):
    run = para.add_run(text)
    run.bold   = bold
    run.italic = italic
    run.font.size = Pt(size)
    run.font.name = 'Garamond'
    if color:
        run.font.color.rgb = RGBColor(*color)
    return run

def add_section_heading(doc, title):
    p = doc.add_paragraph()
    set_spacing(p, before=6, after=1)
    run = p.add_run(title.upper())
    run.bold = True
    run.font.size = Pt(11)
    run.font.name = 'Garamond'
    # Bottom border (rule line)
    pPr = p._element.get_or_add_pPr()
    pBdr = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'),   'single')
    bottom.set(qn('w:sz'),    '6')
    bottom.set(qn('w:space'), '1')
    bottom.set(qn('w:color'), '000000')
    pBdr.append(bottom)
    pPr.append(pBdr)
    return p

def add_subheading(doc, left_bold, right, left_sub_italic, right_sub_italic):
    # Row 1: bold left, right aligned date
    p = doc.add_paragraph()
    set_spacing(p, before=3, after=0)
    tab_stops = p.paragraph_format.tab_stops
    tab_stops.add_tab_stop(Inches(6.5), 2)  # right align
    add_run(p, left_bold, bold=True)
    p.add_run('\t')
    add_run(p, right)
    # Row 2: italic sub
    p2 = doc.add_paragraph()
    set_spacing(p2, before=0, after=0)
    tab_stops2 = p2.paragraph_format.tab_stops
    tab_stops2.add_tab_stop(Inches(6.5), 2)
    add_run(p2, left_sub_italic, italic=True)
    p2.add_run('\t')
    add_run(p2, right_sub_italic, italic=True)

def add_bullet(doc, text):
    p = doc.add_paragraph(style='List Bullet')
    set_spacing(p, before=0, after=1)
    p.paragraph_format.left_indent  = Inches(0.25)
    p.paragraph_format.first_line_indent = Inches(-0.15)
    run = p.add_run(text)
    run.font.size = Pt(9.5)
    run.font.name = 'Garamond'
    return p

def add_skills_line(doc, label, value):
    p = doc.add_paragraph()
    set_spacing(p, before=1, after=1)
    add_run(p, label, bold=True, size=9.5)
    add_run(p, value, size=9.5)

# ============================================================
# HEADING
# ============================================================
name_p = doc.add_paragraph()
name_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
set_spacing(name_p, before=0, after=2)
name_run = name_p.add_run('RICHARD EFFAH YEBOAH')
name_run.bold = True
name_run.font.size = Pt(20)
name_run.font.name = 'Garamond'

contact_p = doc.add_paragraph()
contact_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
set_spacing(contact_p, before=0, after=4)
add_run(contact_p,
    '(614) 290-4260  |  kkras5050@gmail.com  |  github.com/richardeyeboah  |  '
    'linkedin.com/in/richyeff',
    size=9.5)

# ============================================================
# EDUCATION
# ============================================================
add_section_heading(doc, 'Education')
add_subheading(doc,
    'Grambling State University', 'Expected May 2028',
    'Bachelor of Science in Computer Science  |  GPA: 4.0/4.0', 'Grambling, Louisiana')

p = doc.add_paragraph()
set_spacing(p, before=4, after=0)
add_run(p, 'Honors: ', bold=True, size=9.5)
add_run(p, "President's List", size=9.5)

# ============================================================
# EXPERIENCE
# ============================================================
add_section_heading(doc, 'Experience')

add_subheading(doc, 'Kintampo African Market  |  kintampoafricanmarket.com', 'Mar 2026 – Present',
               'Freelance Full-Stack Developer', 'Columbus, Ohio')
add_bullet(doc, 'Rebuilt and deployed a Next.js storefront for 150+ African and Caribbean products with customer accounts, Stripe checkout, local pickup, and nationwide shipping')
add_bullet(doc, 'Implemented server-authoritative price validation, Stripe webhooks, unique payment records, refunds, and order-fulfillment workflows with TypeScript, Supabase, and PostgreSQL')
add_bullet(doc, 'Built staff tools for catalog, inventory, customer, and shipping operations; maintain 61 Playwright tests across 13 files covering checkout, tax, refunds, receipts, navigation, and administrative APIs')

# ============================================================
# PROJECTS
# ============================================================
add_section_heading(doc, 'Projects')

p = doc.add_paragraph()
set_spacing(p, before=3, after=0)
tab_stops = p.paragraph_format.tab_stops
tab_stops.add_tab_stop(Inches(6.5), 2)
add_run(p, 'Prince Auto Inventory & POS System', bold=True)
p.add_run('\t')
add_run(p, 'Aug 2026 – Present')

p2 = doc.add_paragraph()
set_spacing(p2, before=0, after=0)
add_run(p2, 'React, TypeScript, Supabase, PostgreSQL, PWA  |  prince-inventory-manager.vercel.app  |  github.com/richardeyeboah/prince-inventory-manager', italic=True, size=9.5)

add_bullet(doc, 'Built a role-based inventory and point-of-sale application for a mechanic shop with shared stock, parts and labor checkout, customer balances, printable receipts, and reports')
add_bullet(doc, 'Implemented PostgreSQL functions that validate stock and atomically record sales, inventory changes, and financial snapshots; made offline retries idempotent to prevent duplicate sales')
add_bullet(doc, 'Added an offline sales queue that syncs after reconnection and owner-controlled voids that restore stock while preserving an audit trail')

p = doc.add_paragraph()
set_spacing(p, before=3, after=0)
tab_stops = p.paragraph_format.tab_stops
tab_stops.add_tab_stop(Inches(6.5), 2)
add_run(p, 'Item7 Food Truck Management System', bold=True)
p.add_run('\t')
add_run(p, 'Jan 2026')

p2 = doc.add_paragraph()
set_spacing(p2, before=0, after=0)
add_run(p2, 'Python, Flask, PostgreSQL, SQLAlchemy  |  foodtruckk.vercel.app  |  github.com/richardeyeboah/project-foodtruck', italic=True, size=9.5)

add_bullet(doc, 'Developed a responsive ordering and operations application with customer checkout, real-time order tracking, staff scheduling, shift tracking, and role-based management tools')
add_bullet(doc, 'Modeled users, menu items, orders, schedules, and shifts in PostgreSQL with SQLAlchemy; added CSRF protection, rate limiting, password hashing, and input validation')

# ============================================================
# TECHNICAL SKILLS
# ============================================================
add_section_heading(doc, 'Technical Skills')
add_skills_line(doc, 'Languages: ', 'TypeScript, JavaScript, Python, Java, SQL, HTML/CSS')
add_skills_line(doc, 'Frameworks & Runtimes: ', 'Next.js, React, Flask, Node.js, Tailwind CSS')
add_skills_line(doc, 'Data & Services: ', 'PostgreSQL, Supabase, Redis, SQLAlchemy, Stripe')
add_skills_line(doc, 'Testing & Tools: ', 'Playwright, Vitest, Git, GitHub Actions, Vercel')
add_skills_line(doc, 'Engineering: ', 'Full-Stack Development, REST APIs, Authentication, Relational Data Modeling')

# ============================================================
out = "resume.docx"
doc.save(out)
print(f"Saved: {out}")
